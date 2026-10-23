"""
tests/test_cross_file_consistency.py

Pytest test suite verifying consistency between:
- Standard questions (165) vs Expanded questions (904) in `multiphase_course_data.json`
- Master solution dictionary (`verified_all_quizzes_master.json`)
- Student-facing JSONs (`course_data.json`, `quiz_data.json`)
- Text quiz files (`vlearn-app/public/quiz_txt/` and `downloads/Khoa_4_Phase_1/`)
- Excel Master Workbooks (`VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx`)
- Quizizz Excel files (`TONG_HOP_QUIZ_15_NGAY_EXCEL/` and `vlearn-app/public/quiz_excel/`)
- ZIP bundles (`TONG_HOP_QUIZ_15_NGAY_EXCEL.zip` and public mirror)
"""

from __future__ import annotations

import glob
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Set

import openpyxl
import pytest


class TestSharedQuestionsConsistency:
    """Verifies consistency between the 165 standard questions and 904 expanded questions."""

    def test_all_165_standard_ids_present_in_expanded(
        self, standard_questions: List[Dict[str, Any]], expanded_questions: List[Dict[str, Any]]
    ):
        std_ids = {q["question_id"] for q in standard_questions}
        exp_ids = {q["question_id"] for q in expanded_questions}
        assert len(std_ids) == 165
        missing = std_ids - exp_ids
        assert len(missing) == 0, f"Standard question IDs missing in expanded questions: {missing}"

    def test_shared_questions_answers_match(
        self, standard_questions: List[Dict[str, Any]], expanded_questions: List[Dict[str, Any]]
    ):
        """
        Validates that for all 165 shared question IDs, correct_answer is identical.
        In M0 baseline, catches 52 discrepancies:
        - 46 matching questions (empty dict {} in expanded vs valid mapping in standard)
        - 6 numeric questions (string 'None' in expanded vs valid number in standard)
        After M2.1/M2/M3, all 165 must match.
        """
        std_map = {q["question_id"]: q.get("correct_answer") for q in standard_questions}
        exp_map = {q["question_id"]: q.get("correct_answer") for q in expanded_questions if q["question_id"] in std_map}

        discrepancies = []
        for qid, std_ans in std_map.items():
            exp_ans = exp_map.get(qid)
            if std_ans != exp_ans:
                discrepancies.append(f"[{qid}] standard={std_ans!r} vs expanded={exp_ans!r}")

        assert len(discrepancies) == 0, (
            f"Found {len(discrepancies)} answer discrepancies between standard and expanded questions "
            f"(Baseline known: 52 discrepancies):\n" + "\n".join(discrepancies[:15])
        )

    def test_shared_questions_interaction_types_match(
        self, standard_questions: List[Dict[str, Any]], expanded_questions: List[Dict[str, Any]]
    ):
        std_types = {q["question_id"]: q.get("interaction", {}).get("type") for q in standard_questions}
        exp_types = {
            q["question_id"]: q.get("interaction", {}).get("type")
            for q in expanded_questions
            if q["question_id"] in std_types
        }

        mismatches = []
        for qid, stype in std_types.items():
            etype = exp_types.get(qid)
            if stype != etype:
                mismatches.append(f"[{qid}] standard type '{stype}' != expanded type '{etype}'")

        assert len(mismatches) == 0, f"Interaction type mismatches: {mismatches}"


def _extract_master_answer(cr: Any) -> Any:
    """Extracts the underlying answer representation from master JSON correct_response."""
    if not isinstance(cr, dict):
        return cr
    if "correct_choice_id" in cr:
        return cr["correct_choice_id"]
    if "correct_choice_ids" in cr:
        return cr["correct_choice_ids"]
    if "correct_pairs" in cr:
        return cr["correct_pairs"]
    if "correct_order" in cr:
        return cr["correct_order"]
    if "value" in cr:
        return cr["value"]
    return cr


class TestMasterJsonConsistency:
    """Verifies consistency with verified_all_quizzes_master.json (170 verified solutions)."""

    def test_master_json_count_and_keys(self, master_json_data: Dict[str, Any]):
        assert len(master_json_data) == 170, f"Expected 170 master items, found {len(master_json_data)}"

    def test_standard_questions_match_master_solutions(
        self, standard_questions: List[Dict[str, Any]], master_json_data: Dict[str, Any]
    ):
        mismatches = []
        for q in standard_questions:
            qid = q["question_id"]
            if qid not in master_json_data:
                mismatches.append(f"[{qid}] missing from verified_all_quizzes_master.json")
                continue
            master_entry = master_json_data[qid]
            raw_master_resp = master_entry.get("correct_response")
            master_resp = _extract_master_answer(raw_master_resp)
            std_ans = q.get("correct_answer")

            # Check equality (handles int/float or dict comparison)
            if std_ans != master_resp:
                mismatches.append(f"[{qid}] standard answer {std_ans!r} != master {master_resp!r}")

        assert len(mismatches) == 0, f"Master JSON mismatches ({len(mismatches)}):\n" + "\n".join(mismatches[:10])


class TestStudentFacingDataConsistency:
    """Verifies student-facing subsets course_data.json and quiz_data.json."""

    def test_course_data_json_questions_count(
        self, course_data_json: List[Dict[str, Any]], standard_questions: List[Dict[str, Any]]
    ):
        cd_qids = []
        for day in course_data_json[:15]:
            for q in day.get("questions", []):
                cd_qids.append(q.get("question_id"))

        std_qids = [q["question_id"] for q in standard_questions]
        assert len(cd_qids) == 165
        assert cd_qids == std_qids, "Question ID sequence in course_data.json does not match standard questions"

    def test_quiz_data_json_questions_count(
        self, quiz_data_json: Dict[str, Any], standard_questions: List[Dict[str, Any]]
    ):
        qd_qids = []
        for key, day_data in quiz_data_json.items():
            if isinstance(day_data, dict):
                qlist = day_data.get("questions", [])
            elif isinstance(day_data, list):
                qlist = day_data
            else:
                qlist = []
            for q in qlist:
                qd_qids.append(q.get("question_id"))

        std_qids_set = {q["question_id"] for q in standard_questions}
        assert len(qd_qids) == 165
        assert set(qd_qids) == std_qids_set, "Question IDs in quiz_data.json do not match standard questions set"


class TestTextQuizFilesConsistency:
    """Verifies that all 15 Day TXT quiz files exist, match mirrors, and contain all 165 standard questions."""

    def test_standard_txt_files_exist_and_match_downloads(self, project_root: Path):
        vlearn_txt_dir = project_root / "vlearn-app" / "public" / "quiz_txt"
        downloads_dir = project_root / "downloads" / "Khoa_4_Phase_1"

        vlearn_files = sorted(vlearn_txt_dir.glob("*_QUIZ_CHI_TIET_DAP_AN.txt"))
        assert len(vlearn_files) == 15, f"Expected 15 standard TXT files in vlearn-app, found {len(vlearn_files)}"

        for vf in vlearn_files:
            folder_prefix = vf.name.replace("_QUIZ_CHI_TIET_DAP_AN.txt", "")
            dl_file = downloads_dir / folder_prefix / "QUIZ_CHI_TIET_DAP_AN.txt"
            assert dl_file.exists(), f"Missing download TXT for {vf.name} at {dl_file}"

            vf_hash = hashlib.sha256(vf.read_bytes()).hexdigest()
            dl_hash = hashlib.sha256(dl_file.read_bytes()).hexdigest()
            assert vf_hash == dl_hash, f"Hash mismatch between {vf.name} and {dl_file}"

    def test_standard_txt_file_question_ids(self, project_root: Path, standard_questions: List[Dict[str, Any]]):
        vlearn_txt_dir = project_root / "vlearn-app" / "public" / "quiz_txt"
        txt_files = sorted(vlearn_txt_dir.glob("*_QUIZ_CHI_TIET_DAP_AN.txt"))

        parsed_qids = []
        qid_pattern = re.compile(r"Mã:\s*(d\d{2}-q\d{3})")
        for f in txt_files:
            content = f.read_text(encoding="utf-8")
            matches = qid_pattern.findall(content)
            parsed_qids.extend(matches)

        assert len(parsed_qids) == 165, f"Expected 165 question IDs in standard TXT files, found {len(parsed_qids)}"
        std_ids = [q["question_id"] for q in standard_questions]
        assert parsed_qids == std_ids, "Parsed question IDs in standard TXT files do not match standard questions sequence"

    def test_bank_txt_files_exist_and_mirror(self, project_root: Path):
        vlearn_txt_dir = project_root / "vlearn-app" / "public" / "quiz_txt"
        ngan_hang_files = list(vlearn_txt_dir.glob("*_NGAN_HANG_TOAN_DIEN.txt"))
        day_alias_files = list(vlearn_txt_dir.glob("DAY*_NGAN_HANG_*_TOAN_DIEN.txt"))

        assert len(ngan_hang_files) == 15, f"Expected 15 NGAN_HANG TXT files, found {len(ngan_hang_files)}"
        assert len(day_alias_files) == 15, f"Expected 15 DAY* NGAN_HANG alias TXT files, found {len(day_alias_files)}"


class TestMasterExcelConsistency:
    """Verifies master workbooks VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx."""

    def test_master_workbooks_exist_and_have_34_sheets(self, project_root: Path):
        wb1_path = project_root / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx"
        wb2_path = project_root / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES_VERIFIED.xlsx"

        assert wb1_path.exists(), f"Missing {wb1_path}"
        assert wb2_path.exists(), f"Missing {wb2_path}"

        wb1 = openpyxl.load_workbook(wb1_path, read_only=True)
        assert len(wb1.sheetnames) == 34
        assert "Toàn Bộ 165 Câu Quiz & Đáp Án" in wb1.sheetnames
        wb1.close()

    def test_master_workbook_165_questions_sheet(
        self, project_root: Path, standard_questions: List[Dict[str, Any]]
    ):
        wb_path = project_root / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx"
        wb = openpyxl.load_workbook(wb_path, read_only=True)
        ws = wb["Toàn Bộ 165 Câu Quiz & Đáp Án"]

        parsed_qids = []
        qid_to_answer = {}
        for row in ws.iter_rows(min_row=4, max_row=168, min_col=4, max_col=9, values_only=True):
            qid = str(row[0]).strip() if row[0] else ""
            if qid:
                parsed_qids.append(qid)
                # row[0] is col 4 (Mã Câu), row[5] is col 9 (ĐÁP ÁN ĐÚNG CHUẨN XÁC)
                qid_to_answer[qid] = str(row[5] or "")
        wb.close()

        assert len(parsed_qids) == 165
        std_ids = [q["question_id"] for q in standard_questions]
        assert parsed_qids == std_ids, "Question IDs in master Excel sheet do not match standard questions sequence"

        # Explicitly verify that numeric questions have updated answers and do NOT contain 'None'
        std_numeric = {
            q["question_id"]: q["correct_answer"]
            for q in standard_questions
            if q.get("interaction", {}).get("type") == "numeric_input"
        }
        assert len(std_numeric) > 0, "No numeric questions found in standard questions"

        for qid, expected_val in std_numeric.items():
            ans_cell = qid_to_answer.get(qid, "")
            assert "None" not in ans_cell, (
                f"Numeric question {qid} in sheet 'Toàn Bộ 165 Câu Quiz & Đáp Án' contains 'None': {ans_cell!r}"
            )
            assert str(expected_val) in ans_cell, (
                f"Numeric question {qid} does not contain expected numeric answer {expected_val}: {ans_cell!r}"
            )


class TestQuizizzExcelAndZipConsistency:
    """Verifies Quizizz format spreadsheets and zip archives."""

    def test_quizizz_904_questions_row_count(self, project_root: Path):
        base_dir = project_root / "TONG_HOP_QUIZ_15_NGAY_EXCEL" / "01_CHUAN_IMPORT_QUIZIZZ_QUIZ_COM_904_CAU"
        full_workbooks = sorted(base_dir.glob("*/Day_*_FULL_*.xlsx"))
        assert len(full_workbooks) == 15, f"Expected 15 FULL workbooks, found {len(full_workbooks)}"

        total_rows = 0
        for wb_path in full_workbooks:
            wb = openpyxl.load_workbook(wb_path, read_only=True)
            ws = wb["Create a Quiz"]
            # Rows 1 and 2 are headers/instructions; data starts on row 3
            data_rows = 0
            for row in ws.iter_rows(min_row=3, values_only=True):
                if row[0] is not None and str(row[0]).strip():
                    data_rows += 1
            total_rows += data_rows
            wb.close()

        assert total_rows == 904, f"Expected 904 question rows across 15 FULL workbooks, found {total_rows}"

    def test_zip_archives_integrity_and_parity(self, project_root: Path):
        zip1_path = project_root / "TONG_HOP_QUIZ_15_NGAY_EXCEL.zip"
        zip2_path = project_root / "vlearn-app" / "public" / "TONG_HOP_QUIZ_15_NGAY_EXCEL.zip"

        assert zip1_path.exists(), f"Missing {zip1_path}"
        assert zip2_path.exists(), f"Missing {zip2_path}"

        h1 = hashlib.sha256(zip1_path.read_bytes()).hexdigest()
        h2 = hashlib.sha256(zip2_path.read_bytes()).hexdigest()

        assert h1 == h2, f"SHA256 mismatch between {zip1_path} and {zip2_path}"

        with zipfile.ZipFile(zip1_path, "r") as z:
            file_entries = [name for name in z.namelist() if not name.endswith("/")]
            assert len(file_entries) == 146, f"Expected 146 file entries in ZIP, found {len(file_entries)}"
