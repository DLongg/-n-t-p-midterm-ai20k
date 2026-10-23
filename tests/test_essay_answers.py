"""
tests/test_essay_answers.py

Comprehensive End-to-End (E2E) automated test suite and validation engine for
the 775 essay questions across Days 01–15 of the VLearn learning platform.

Validation Architecture:
- Tier 1: Feature Coverage (775 questions, day breakdown, interaction type, required fields)
- Tier 2: Boundary & Corner Cases (no boilerplate, explanation != hint, answer len >= 150, points match rubric, reference source)
- Tier 3: Cross-File & Sync Parity (Master JSON vs Bank TXT, Master Excel, Partitioned Excel, ZIP archives)
- Tier 4: Application Scenarios (Vite production build and App.jsx UI integration)
- Chunk Files Validation: Standalone CLI and test runner for `data/essay_block_*.json`

Author: Test Writer (E2E Testing Track)
Milestone: M1 / M4
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent

# ---------------------------------------------------------------------------
# Canonical Constants & Day-by-Day Question Matrix
# ---------------------------------------------------------------------------

TOTAL_EXPECTED_ESSAY = 775

EXPECTED_DAY_ESSAY_COUNTS: dict[int, int] = {
    1: 49,
    2: 86,
    3: 52,
    4: 30,
    5: 38,
    6: 27,
    7: 45,
    8: 55,
    9: 48,
    10: 48,
    11: 57,
    12: 73,
    13: 61,
    14: 52,
    15: 54,
}

BLOCK_METADATA: dict[str, dict[str, Any]] = {
    "essay_block_01_03.json": {
        "days": [1, 2, 3],
        "expected_count": 187,
        "description": "Days 01–03 Model Answers & Rubrics (49 + 86 + 52)",
    },
    "essay_block_04_06.json": {
        "days": [4, 5, 6],
        "expected_count": 95,
        "description": "Days 04–06 Model Answers & Rubrics (30 + 38 + 27)",
    },
    "essay_block_07_09.json": {
        "days": [7, 8, 9],
        "expected_count": 148,
        "description": "Days 07–09 Model Answers & Rubrics (45 + 55 + 48)",
    },
    "essay_block_10_12.json": {
        "days": [10, 11, 12],
        "expected_count": 178,
        "description": "Days 10–12 Model Answers & Rubrics (48 + 57 + 73)",
    },
    "essay_block_13_15.json": {
        "days": [13, 14, 15],
        "expected_count": 167,
        "description": "Days 13–15 Model Answers & Rubrics (61 + 52 + 54)",
    },
}

BOILERPLATE_PATTERNS = [
    "Đúng bản chất vấn đề và yêu cầu kỹ thuật",
    "Đầy đủ các tiêu chí chấp nhận / test case tối thiểu",
    "Phân tích theo nguyên lý thiết kế: Với mỗi mũi tên, hỏi",
    "[TODO]",
    "[TBD]",
    "Chưa có đáp án",
    "Cập nhật sau",
]

MIN_SAMPLE_ANSWER_LEN = 150
MIN_EXPLANATION_LEN = 50
MIN_REFERENCE_SOURCE_LEN = 10

# ---------------------------------------------------------------------------
# Validation Logic & Helpers (Usable by Pytest and as Standalone CLI/Library)
# ---------------------------------------------------------------------------

def parse_rubric_points(rubric: str, expected_points: int | None) -> list[str]:
    """
    Parses rubric points and checks consistency with expected_points.
    Checks:
    1. Stated points in header: e.g., '👉 Tiêu chí chấm (Rubric 3 điểm):' or 'Rubric 3 điểm'
    2. Sum of individual criterion sub-points: e.g., '- Tiêu chí 1 (1.5 điểm)'
    """
    if not rubric or not isinstance(rubric, str):
        return ["Rubric is empty or not a string"]

    errors = []
    lines = rubric.splitlines()

    # 1. Header points inspection
    header_m = re.search(
        r"(?:Rubric|Thang điểm|Điểm tối đa)\s*:?\s*(\d+(?:\.\d+)?)\s*(?:điểm|đ)",
        rubric,
        re.IGNORECASE,
    )
    header_pts: float | None = float(header_m.group(1)) if header_m else None

    if expected_points is not None and header_pts is not None and abs(header_pts - expected_points) > 0.01:
        errors.append(
            f"Rubric header points ({header_pts}) does not match question points ({expected_points})"
        )

    # 2. Sub-criteria points inspection (excluding the header line itself)
    sub_pts: list[float] = []
    for line in lines:
        if header_m and line.strip() == header_m.group(0).strip():
            continue
        # Check bulleted or numbered lines with point allocations
        # Prioritize explicit parenthesized/bracketed points like (1.0 điểm)
        paren_m = re.search(r"[\(\[]\s*(\d+(?:\.\d+)?)\s*(?:điểm|đ)\s*[\)\]]", line, re.IGNORECASE)
        if paren_m:
            try:
                sub_pts.append(float(paren_m.group(1)))
                continue
            except ValueError:
                pass

        bullet_m = re.search(
            r"(?:^|\s)(?:[-*•]|\d+\.)\s*.*?(?:[\(\[]|\b:?\s*)(\d+(?:\.\d+)?)\s*(?:điểm|đ)\b",
            line,
            re.IGNORECASE,
        )
        if bullet_m:
            try:
                sub_pts.append(float(bullet_m.group(1)))
            except ValueError:
                pass

    if sub_pts and expected_points is not None:
        sub_sum = sum(sub_pts)
        if abs(sub_sum - expected_points) > 0.05:
            errors.append(
                f"Sum of rubric criteria points ({sub_sum}) does not match question points ({expected_points})"
            )
    elif header_pts is None and expected_points is not None:
        errors.append(
            f"No recognizable point values found in rubric (expected {expected_points} điểm)"
        )

    return errors


def validate_essay_entry(
    qid: str,
    data: dict[str, Any],
    expected_points: int | None = None,
    hints: list[dict[str, Any]] | None = None,
) -> list[str]:
    """
    Validates a single essay question entry against all quality criteria:
    - Required fields: sample_answer, explanation, rubric, reference_source
    - Minimum lengths: sample_answer >= 150 chars, explanation >= 50, reference >= 10
    - Absence of generic boilerplate
    - Explanation must not merely repeat hints
    - Points match rubric stated total/criteria
    """
    errors: list[str] = []

    # Extract fields (supports both chunk entry dict and master question dict)
    sample_answer = data.get("sample_answer")
    explanation = data.get("explanation")
    rubric = data.get("rubric")
    reference_source = data.get("reference_source")

    pts = data.get("points") if "points" in data else expected_points
    q_hints = data.get("hints") if "hints" in data else hints

    # 1. Field presence and types
    if not sample_answer or not isinstance(sample_answer, str):
        errors.append("Missing or invalid 'sample_answer'")
    if not explanation or not isinstance(explanation, str):
        errors.append("Missing or invalid 'explanation'")
    if not rubric or not isinstance(rubric, str):
        errors.append("Missing or invalid 'rubric'")
    if not reference_source or not isinstance(reference_source, str):
        errors.append("Missing or invalid 'reference_source'")

    # If any required field is missing, return early
    if errors:
        return errors

    # 2. Length constraints
    if len(sample_answer.strip()) < MIN_SAMPLE_ANSWER_LEN:
        errors.append(
            f"sample_answer too short ({len(sample_answer.strip())} chars < {MIN_SAMPLE_ANSWER_LEN})"
        )

    if len(explanation.strip()) < MIN_EXPLANATION_LEN:
        errors.append(
            f"explanation too short ({len(explanation.strip())} chars < {MIN_EXPLANATION_LEN})"
        )

    if len(reference_source.strip()) < MIN_REFERENCE_SOURCE_LEN:
        errors.append(
            f"reference_source too short ({len(reference_source.strip())} chars < {MIN_REFERENCE_SOURCE_LEN})"
        )

    # 3. Boilerplate checks
    for bp in BOILERPLATE_PATTERNS:
        if bp.lower() in rubric.lower():
            errors.append(f"Rubric retains generic boilerplate: '{bp}'")
        if bp.lower() in explanation.lower():
            errors.append(f"Explanation retains generic boilerplate: '{bp}'")
        if bp.lower() in sample_answer.lower():
            errors.append(f"Sample answer retains generic boilerplate: '{bp}'")

    # 4. Explanation distinct from hint
    if q_hints and isinstance(q_hints, list):
        clean_exp = explanation.strip().lower()
        for h in q_hints:
            if isinstance(h, dict) and "text" in h:
                h_text = h["text"].replace("💡", "").strip().lower()
                if not h_text:
                    continue
                # Direct match
                if clean_exp == h_text:
                    errors.append("Explanation is identical to hint text")
                # Prepending 'phân tích theo nguyên lý thiết kế:' wrapper
                elif clean_exp.replace("phân tích theo nguyên lý thiết kế:", "").strip() == h_text:
                    errors.append(
                        "Explanation is merely hint text prefixed with 'Phân tích theo nguyên lý thiết kế:'"
                    )
                # Trivial wrapper of hint with no substantive answer
                elif h_text in clean_exp and len(clean_exp) < len(h_text) + 40:
                    errors.append("Explanation is a trivial wrapper around the hint without substantive analysis")

    # 5. Points match rubric
    if pts is not None:
        pts_errors = parse_rubric_points(rubric, int(pts))
        errors.extend(pts_errors)

    return errors


def load_master_metadata(project_root: Path) -> dict[str, dict[str, Any]]:
    """Loads metadata index for all 775 essay questions from multiphase_course_data.json."""
    data_path = project_root / "vlearn-app" / "src" / "multiphase_course_data.json"
    if not data_path.exists():
        raise FileNotFoundError(f"Missing master data: {data_path}")

    with open(data_path, "r", encoding="utf-8") as f:
        multiphase = json.load(f)

    days = multiphase.get("phases", [])[0].get("days", [])[:15]
    metadata: dict[str, dict[str, Any]] = {}
    for day_idx, day in enumerate(days, 1):
        for q in day.get("essay_questions", []):
            qid = q["question_id"]
            metadata[qid] = {
                "day_num": day_idx,
                "points": q.get("points"),
                "hints": q.get("hints", []),
                "prompt": q.get("interaction", {}).get("prompt", ""),
            }
    return metadata


def validate_chunk_file(chunk_path: Path, master_metadata: dict[str, dict[str, Any]] | None = None) -> dict[str, Any]:
    """
    Validates an individual authoring chunk file (e.g. data/essay_block_01_03.json).
    Returns a diagnostic report dictionary.
    """
    if not chunk_path.exists():
        return {
            "chunk": chunk_path.name,
            "status": "MISSING",
            "error": f"File does not exist: {chunk_path}",
        }

    with open(chunk_path, "r", encoding="utf-8") as f:
        try:
            chunk_data = json.load(f)
        except (json.JSONDecodeError, OSError, ValueError) as e:
            return {
                "chunk": chunk_path.name,
                "status": "INVALID_JSON",
                "error": f"JSON parse error: {e}",
            }

    fname = chunk_path.name
    meta_info = BLOCK_METADATA.get(fname, {})
    expected_count = meta_info.get("expected_count")

    if master_metadata is None:
        master_metadata = load_master_metadata(PROJECT_ROOT)

    chunk_errors: dict[str, list[str]] = {}
    valid_count = 0

    for qid, qdata in chunk_data.items():
        qmeta = master_metadata.get(qid, {})
        pts = qmeta.get("points")
        hints = qmeta.get("hints")
        errs = validate_essay_entry(qid, qdata, expected_points=pts, hints=hints)
        if errs:
            chunk_errors[qid] = errs
        else:
            valid_count += 1

    # Check for missing questions if expected count is known
    missing_qids: list[str] = []
    if meta_info:
        target_days = meta_info.get("days", [])
        expected_qids = [
            qid for qid, info in master_metadata.items() if info["day_num"] in target_days
        ]
        missing_qids = [qid for qid in expected_qids if qid not in chunk_data]

    total_in_chunk = len(chunk_data)
    is_clean = len(chunk_errors) == 0 and len(missing_qids) == 0
    if expected_count is not None and total_in_chunk != expected_count:
        is_clean = False

    return {
        "chunk": fname,
        "status": "CLEAN" if is_clean else "DEFECTIVE",
        "total_items": total_in_chunk,
        "expected_items": expected_count,
        "valid_items": valid_count,
        "missing_items": missing_qids,
        "item_errors": chunk_errors,
    }

# ---------------------------------------------------------------------------
# Pytest Session Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(scope="session")
def essay_questions_by_day(phase1_days: list[dict[str, Any]]) -> dict[int, list[dict[str, Any]]]:
    """Returns essay questions mapped by Day number (1..15)."""
    mapping: dict[int, list[dict[str, Any]]] = {}
    for day_idx, day in enumerate(phase1_days, 1):
        mapping[day_idx] = day.get("essay_questions", [])
    return mapping


@pytest.fixture(scope="session")
def all_essay_questions(essay_questions_by_day: dict[int, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    """Returns a flat list of all 775 essay questions with _day_num attached."""
    all_qs: list[dict[str, Any]] = []
    for day_num, qs in essay_questions_by_day.items():
        for q in qs:
            item = dict(q)
            item["_day_num"] = day_num
            all_qs.append(item)
    return all_qs

# ---------------------------------------------------------------------------
# Tier 1: Feature Coverage (775 Questions, Day Breakdown, Schema)
# ---------------------------------------------------------------------------

class TestTier1FeatureCoverage:
    """Verifies that each of the 15 Days contains expected counts and valid schemas."""

    def test_total_essay_questions_count(self, all_essay_questions: list[dict[str, Any]]):
        """Total essay questions across Phase 1 Days 01–15 must be exactly 775."""
        assert len(all_essay_questions) == TOTAL_EXPECTED_ESSAY, (
            f"Expected {TOTAL_EXPECTED_ESSAY} total essay questions across Days 01–15, "
            f"found {len(all_essay_questions)}"
        )

    @pytest.mark.parametrize("day_num", list(range(1, 16)))
    def test_day_by_day_essay_counts(
        self, essay_questions_by_day: dict[int, list[dict[str, Any]]], day_num: int
    ):
        """Verifies exact count of essay questions for each specific Day 01–15."""
        expected = EXPECTED_DAY_ESSAY_COUNTS[day_num]
        actual = len(essay_questions_by_day[day_num])
        assert actual == expected, (
            f"Day {day_num:02d}: expected {expected} essay questions, found {actual}"
        )

    def test_all_essay_questions_unique_ids(self, all_essay_questions: list[dict[str, Any]]):
        """Verifies all 775 essay question IDs are unique and conform to standard regex."""
        id_pattern = re.compile(r"^d\d{2}-q\d{3}$")
        all_ids = [q["question_id"] for q in all_essay_questions]
        unique_ids = set(all_ids)

        assert len(all_ids) == len(unique_ids), (
            f"Found duplicate essay question IDs: {len(all_ids)} total vs {len(unique_ids)} unique"
        )

        invalid_ids = [qid for qid in all_ids if not id_pattern.match(qid)]
        assert len(invalid_ids) == 0, f"Found malformed question IDs: {invalid_ids[:10]}"

    def test_all_essay_questions_interaction_type_and_prompt(
        self, all_essay_questions: list[dict[str, Any]]
    ):
        """Verifies interaction type is 'essay_short_text' and prompt is non-empty."""
        defects = []
        for q in all_essay_questions:
            qid = q.get("question_id")
            interaction = q.get("interaction", {})
            itype = interaction.get("type")
            prompt = interaction.get("prompt", "")

            if itype != "essay_short_text":
                defects.append(f"[{qid}] invalid interaction type: {itype!r}")
            if not prompt or len(prompt.strip()) < 10:
                defects.append(f"[{qid}] prompt is empty or too short (< 10 chars)")

        assert len(defects) == 0, (
            f"Found {len(defects)} interaction defects in essay questions:\n"
            + "\n".join(defects[:10])
        )

    def test_all_essay_questions_have_required_fields(
        self, all_essay_questions: list[dict[str, Any]]
    ):
        """
        Verifies that all 775 essay questions contain required fields:
        question_id, points, rubric, explanation, sample_answer, reference_source.
        """
        required_fields = ["question_id", "points", "rubric", "explanation", "sample_answer", "reference_source"]
        missing_records = []

        for q in all_essay_questions:
            qid = q.get("question_id")
            missing = [f for f in required_fields if f not in q or q[f] is None]
            if missing:
                missing_records.append(f"[{qid}] missing fields: {missing}")

        assert len(missing_records) == 0, (
            f"Found {len(missing_records)} essay questions with missing required fields "
            f"(Baseline expected: missing sample_answer & reference_source until M2/M3 merge):\n"
            + "\n".join(missing_records[:10])
        )

# ---------------------------------------------------------------------------
# Tier 2: Boundary & Corner Cases (Boilerplate, Length, Points Match, Source)
# ---------------------------------------------------------------------------

class TestTier2BoundaryAndCornerCases:
    """Verifies content quality, absence of boilerplate, length, and technical attribution."""

    def test_no_boilerplate_in_rubrics(self, all_essay_questions: list[dict[str, Any]]):
        """Verifies zero rubrics retain generic boilerplate strings or placeholder markers."""
        boilerplate_hits = []
        for q in all_essay_questions:
            qid = q.get("question_id")
            rubric = q.get("rubric", "")
            if not rubric:
                continue
            for bp in BOILERPLATE_PATTERNS:
                if bp.lower() in rubric.lower():
                    boilerplate_hits.append(f"[{qid}] rubric retains boilerplate: {bp!r}")
                    break

        assert len(boilerplate_hits) == 0, (
            f"Found {len(boilerplate_hits)} questions with boilerplate rubrics "
            f"(Baseline expected: 775 before authoring):\n"
            + "\n".join(boilerplate_hits[:10])
        )

    def test_no_boilerplate_or_hint_repeating_in_explanations(
        self, all_essay_questions: list[dict[str, Any]]
    ):
        """
        Verifies explanations are substantive (>= 50 chars), not generic boilerplate,
        and not merely identical or trivial wrappers of the 1-sentence hint.
        """
        defects = []
        for q in all_essay_questions:
            qid = q.get("question_id")
            exp = q.get("explanation", "").strip()
            hints = q.get("hints", [])

            if len(exp) < MIN_EXPLANATION_LEN:
                defects.append(f"[{qid}] explanation too short ({len(exp)} < {MIN_EXPLANATION_LEN})")
                continue

            for bp in BOILERPLATE_PATTERNS:
                if bp.lower() in exp.lower():
                    defects.append(f"[{qid}] explanation retains boilerplate: {bp!r}")
                    break

            # Check hint repetition
            clean_exp = exp.lower()
            for h in hints:
                if isinstance(h, dict) and "text" in h:
                    htext = h["text"].replace("💡", "").strip().lower()
                    if not htext:
                        continue
                    if clean_exp == htext:
                        defects.append(f"[{qid}] explanation is identical to hint")
                    elif clean_exp.replace("phân tích theo nguyên lý thiết kế:", "").strip() == htext:
                        defects.append(
                            f"[{qid}] explanation merely prepends 'Phân tích theo nguyên lý thiết kế:' to hint"
                        )
                    elif htext in clean_exp and len(clean_exp) < len(htext) + 40:
                        defects.append(f"[{qid}] explanation is a trivial wrapper around hint")

        assert len(defects) == 0, (
            f"Found {len(defects)} defective explanations "
            f"(Baseline expected: 775 before authoring):\n"
            + "\n".join(defects[:10])
        )

    def test_sample_answer_minimum_length(self, all_essay_questions: list[dict[str, Any]]):
        """Verifies every sample_answer exists and has length >= 150 characters."""
        defects = []
        for q in all_essay_questions:
            qid = q.get("question_id")
            ans = q.get("sample_answer")
            if not ans or not isinstance(ans, str):
                defects.append(f"[{qid}] missing sample_answer")
            elif len(ans.strip()) < MIN_SAMPLE_ANSWER_LEN:
                defects.append(
                    f"[{qid}] sample_answer too short ({len(ans.strip())} < {MIN_SAMPLE_ANSWER_LEN})"
                )

        assert len(defects) == 0, (
            f"Found {len(defects)} questions failing sample_answer minimum length "
            f"(Baseline expected: 775 before authoring):\n"
            + "\n".join(defects[:10])
        )

    def test_points_match_rubric_total(self, all_essay_questions: list[dict[str, Any]]):
        """
        Verifies that for every question, points stated in rubric header and/or
        sum of criterion sub-points matches q.points.
        """
        discrepancies = []
        for q in all_essay_questions:
            qid = q.get("question_id")
            pts = q.get("points")
            rubric = q.get("rubric", "")

            errs = parse_rubric_points(rubric, pts)
            if errs:
                discrepancies.append(f"[{qid}] points={pts}: {'; '.join(errs)}")

        assert len(discrepancies) == 0, (
            f"Found {len(discrepancies)} point mismatches in rubrics:\n"
            + "\n".join(discrepancies[:10])
        )

    def test_reference_source_validity(self, all_essay_questions: list[dict[str, Any]]):
        """Verifies every question has a valid reference_source (>= 10 chars)."""
        invalid_sources = []
        for q in all_essay_questions:
            qid = q.get("question_id")
            src = q.get("reference_source")
            if not src or not isinstance(src, str) or len(src.strip()) < MIN_REFERENCE_SOURCE_LEN:
                invalid_sources.append(f"[{qid}] invalid reference_source: {src!r}")

        assert len(invalid_sources) == 0, (
            f"Found {len(invalid_sources)} questions with invalid reference_source "
            f"(Baseline expected: 775 before authoring):\n"
            + "\n".join(invalid_sources[:10])
        )

# ---------------------------------------------------------------------------
# Tier 3: Cross-File & Sync Parity (TXT, Master Excel, Partitioned Excel, ZIP)
# ---------------------------------------------------------------------------

class TestTier3CrossFileAndSyncParity:
    """Verifies synchronization across all project distribution formats."""

    def test_bank_txt_files_exist_and_cover_775_questions(self, project_root: Path):
        """Verifies 15 Bank TXT files and 15 alias files exist and contain all 775 question IDs."""
        txt_dir = project_root / "vlearn-app" / "public" / "quiz_txt"
        bank_files = sorted(txt_dir.glob("*_NGAN_HANG_TOAN_DIEN.txt"))
        alias_files = sorted(txt_dir.glob("DAY*_NGAN_HANG_*_CAU_HOI_TOAN_DIEN.txt"))

        assert len(bank_files) == 15, f"Expected 15 bank TXT files, found {len(bank_files)}"
        assert len(alias_files) == 15, f"Expected 15 alias TXT files, found {len(alias_files)}"

        total_txt_qids = []
        qid_pat = re.compile(r"\b(d\d{2}-q\d{3})\b")
        for bf in bank_files:
            text = bf.read_text(encoding="utf-8")
            total_txt_qids.extend(qid_pat.findall(text))

        # Check that all 775 essay question IDs appear in bank TXT files
        master_metadata = load_master_metadata(project_root)
        missing_in_txt = [qid for qid in master_metadata if qid not in total_txt_qids]
        assert len(missing_in_txt) == 0, f"Essay questions missing in Bank TXT files: {missing_in_txt[:10]}"

    def test_bank_txt_files_contain_updated_essay_content(self, project_root: Path):
        """Verifies that across the 15 bank TXT files, zero essay questions retain boilerplate."""
        txt_dir = project_root / "vlearn-app" / "public" / "quiz_txt"
        bank_files = sorted(txt_dir.glob("*_NGAN_HANG_TOAN_DIEN.txt"))

        boilerplate_matches = []
        for bf in bank_files:
            text = bf.read_text(encoding="utf-8")
            if "Đúng bản chất vấn đề và yêu cầu kỹ thuật" in text:
                boilerplate_matches.append(bf.name)

        assert len(boilerplate_matches) == 0, (
            f"Bank TXT files retain boilerplate string 'Đúng bản chất vấn đề...': {boilerplate_matches} "
            f"(Baseline expected until M3.2 sync)"
        )

    def test_master_excel_workbooks_exist_and_cover_775_questions(self, project_root: Path):
        """Verifies master workbooks have 15 Kho sheets containing exactly 775 essay questions."""
        wb_paths = [
            project_root / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx",
            project_root / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES_VERIFIED.xlsx",
        ]
        for p in wb_paths:
            assert p.exists(), f"Missing workbook: {p}"

        wb = openpyxl.load_workbook(wb_paths[0], read_only=True)
        kho_sheets = [s for s in wb.sheetnames if "_Kho_" in s]
        assert len(kho_sheets) == 15, f"Expected 15 Kho sheets, found {len(kho_sheets)}"

        total_essay_rows = 0
        for s in kho_sheets:
            ws = wb[s]
            for r in ws.iter_rows(min_row=4, min_col=4, max_col=4, values_only=True):
                if r[0] and "Tự luận" in str(r[0]):
                    total_essay_rows += 1
        wb.close()

        assert total_essay_rows == TOTAL_EXPECTED_ESSAY, (
            f"Expected {TOTAL_EXPECTED_ESSAY} essay questions across Kho sheets, found {total_essay_rows}"
        )

    def test_master_excel_workbooks_contain_updated_essay_content(self, project_root: Path):
        """Verifies that master workbooks contain updated rubrics without boilerplate."""
        wb_path = project_root / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx"
        wb = openpyxl.load_workbook(wb_path, read_only=True)
        kho_sheets = [s for s in wb.sheetnames if "_Kho_" in s]

        boilerplate_cells = []
        for s in kho_sheets:
            ws = wb[s]
            for r in ws.iter_rows(min_row=4, min_col=2, max_col=9, values_only=True):
                qid, qtype, rubric = r[0], r[2], r[7]
                if qtype and "Tự luận" in str(qtype) and rubric and "Đúng bản chất vấn đề và yêu cầu kỹ thuật" in str(rubric):
                    boilerplate_cells.append(f"[{qid} in {s}] retains boilerplate rubric")
        wb.close()

        assert len(boilerplate_cells) == 0, (
            f"Found {len(boilerplate_cells)} essay rows with boilerplate in Master Excel "
            f"(Baseline expected: 775 until M3.2 sync):\n"
            + "\n".join(boilerplate_cells[:10])
        )

    def test_partitioned_excel_workbooks_exist_and_cover_775_questions(self, project_root: Path):
        """Verifies 84 partitioned workbooks contain all 775 Open-Ended rows."""
        part_dir = (
            project_root
            / "TONG_HOP_QUIZ_15_NGAY_EXCEL"
            / "02_FULL_NGAN_HANG_1679_CAU_CHI_TIET_CHIA_20_CAU"
        )
        xlsx_files = list(part_dir.glob("*/*.xlsx"))
        assert len(xlsx_files) == 84, f"Expected 84 partitioned Excel files, found {len(xlsx_files)}"

        total_open_ended = 0
        for f in xlsx_files:
            wb = openpyxl.load_workbook(f, read_only=True)
            ws = wb["Create a Quiz"]
            for r in ws.iter_rows(min_row=3, min_col=2, max_col=2, values_only=True):
                if r[0] == "Open-Ended":
                    total_open_ended += 1
            wb.close()

        assert total_open_ended == TOTAL_EXPECTED_ESSAY, (
            f"Expected {TOTAL_EXPECTED_ESSAY} Open-Ended questions across 84 partitioned workbooks, "
            f"found {total_open_ended}"
        )

    def test_partitioned_excel_workbooks_contain_updated_essay_content(self, project_root: Path):
        """Verifies that Open-Ended questions in 84 partitioned workbooks do not contain boilerplate."""
        part_dir = (
            project_root
            / "TONG_HOP_QUIZ_15_NGAY_EXCEL"
            / "02_FULL_NGAN_HANG_1679_CAU_CHI_TIET_CHIA_20_CAU"
        )
        xlsx_files = list(part_dir.glob("*/*.xlsx"))

        boilerplate_count = 0
        sample_defects = []
        for f in xlsx_files:
            wb = openpyxl.load_workbook(f, read_only=True)
            ws = wb["Create a Quiz"]
            for r in ws.iter_rows(min_row=3, min_col=2, max_col=11, values_only=True):
                qtype, exp = r[0], r[9]
                if qtype == "Open-Ended" and exp and "Đúng bản chất vấn đề và yêu cầu kỹ thuật" in str(exp):
                    boilerplate_count += 1
                    if len(sample_defects) < 5:
                        sample_defects.append(f"{f.name}: retains boilerplate")
            wb.close()

        assert boilerplate_count == 0, (
            f"Found {boilerplate_count} Open-Ended rows retaining boilerplate in partitioned workbooks "
            f"(Baseline expected: 775 until M3.2 sync):\n"
            + "\n".join(sample_defects)
        )

    def test_zip_archives_parity_and_count(self, project_root: Path):
        """Verifies root and public ZIP archives exist, match bit-for-bit SHA256, and have 146 files."""
        zip1 = project_root / "TONG_HOP_QUIZ_15_NGAY_EXCEL.zip"
        zip2 = project_root / "vlearn-app" / "public" / "TONG_HOP_QUIZ_15_NGAY_EXCEL.zip"

        assert zip1.exists(), f"Missing {zip1}"
        assert zip2.exists(), f"Missing {zip2}"

        h1 = hashlib.sha256(zip1.read_bytes()).hexdigest()
        h2 = hashlib.sha256(zip2.read_bytes()).hexdigest()
        assert h1 == h2, f"SHA256 mismatch between zip archives: {h1} != {h2}"

        with zipfile.ZipFile(zip1, "r") as z:
            entries = [n for n in z.namelist() if not n.endswith("/")]
            assert len(entries) == 146, f"Expected 146 entries in ZIP, found {len(entries)}"

# ---------------------------------------------------------------------------
# Tier 4: Application Scenarios (Build and Web UI Integration)
# ---------------------------------------------------------------------------

class TestTier4ApplicationScenarios:
    """Verifies that the web application builds cleanly and UI integrates essay questions."""

    def test_npm_run_build_production_bundle(self, project_root: Path):
        """Verifies `npm run build` succeeds with exit code 0 and outputs production artifacts."""
        vlearn_dir = project_root / "vlearn-app"
        package_json = vlearn_dir / "package.json"
        assert package_json.exists(), f"package.json missing at {package_json}"

        npm_cmd = "npm.cmd" if sys.platform == "win32" else "npm"
        proc = subprocess.run(
            [npm_cmd, "run", "build"],
            cwd=str(vlearn_dir),
            capture_output=True,
            text=True,
            timeout=120,
            shell=(sys.platform == "win32"),
            check=False,
        )

        assert proc.returncode == 0, (
            f"npm run build failed with exit code {proc.returncode}!\n"
            f"STDOUT:\n{proc.stdout}\n"
            f"STDERR:\n{proc.stderr}"
        )

        dist_dir = vlearn_dir / "dist"
        index_html = dist_dir / "index.html"
        assets_dir = dist_dir / "assets"

        assert dist_dir.exists(), f"dist directory was not created at {dist_dir}"
        assert index_html.exists() and index_html.stat().st_size > 0, "index.html missing or empty"
        assert assets_dir.exists(), f"assets directory missing at {assets_dir}"
        assert len(list(assets_dir.glob("*.js"))) >= 1, "No compiled JS bundle found in assets directory"

    def test_vlearn_app_essay_ui_readiness(self, project_root: Path):
        """Verifies App.jsx contains handlers for essay questions and study mode."""
        app_jsx_path = project_root / "vlearn-app" / "src" / "App.jsx"
        assert app_jsx_path.exists(), f"Missing App.jsx: {app_jsx_path}"

        code = app_jsx_path.read_text(encoding="utf-8")
        assert "essay_short_text" in code, "App.jsx does not reference 'essay_short_text'"
        assert "showEssaySolution" in code, "App.jsx does not manage 'showEssaySolution' state"
        assert "rubric" in code, "App.jsx does not reference rubric"
        assert "studyMode" in code, "App.jsx does not manage studyMode"

# ---------------------------------------------------------------------------
# Dynamic Chunk Files Validation (Authoring Blocks)
# ---------------------------------------------------------------------------

class TestChunkFilesValidation:
    """Dynamically validates authoring chunk files in data/ directory as they become available."""

    def test_available_chunk_files_compliance(self, project_root: Path):
        """
        Discovers any available chunk files in `data/` (e.g. data/essay_block_*.json)
        and verifies that all authored items meet strict quality gates.
        """
        data_dir = project_root / "data"
        if not data_dir.exists():
            pytest.skip("data/ directory does not exist yet (M2 in progress)")

        chunk_files = sorted(data_dir.glob("essay_block_*.json"))
        if not chunk_files:
            pytest.skip("No chunk files found in data/ yet (authoring in progress)")

        master_metadata = load_master_metadata(project_root)
        total_failures = []

        for cf in chunk_files:
            res = validate_chunk_file(cf, master_metadata)
            if res.get("status") != "CLEAN":
                err_summary = []
                if res.get("missing_items"):
                    err_summary.append(f"Missing {len(res['missing_items'])} questions: {res['missing_items'][:5]}")
                if res.get("item_errors"):
                    for qid, errs in list(res["item_errors"].items())[:5]:
                        err_summary.append(f"[{qid}]: {'; '.join(errs)}")
                total_failures.append(f"{cf.name}: " + " | ".join(err_summary))

        assert len(total_failures) == 0, (
            "Chunk file quality violations detected:\n" + "\n".join(total_failures)
        )

# ---------------------------------------------------------------------------
# Standalone CLI Entrypoint
# ---------------------------------------------------------------------------

def main() -> int:
    """CLI runner for standalone chunk and master validation."""
    parser = argparse.ArgumentParser(
        description="VLearn Essay Questions Automated Test & Validation Engine",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--chunk",
        type=str,
        help="Path to a single chunk JSON file (e.g. data/essay_block_01_03.json)",
    )
    parser.add_argument(
        "--all-chunks",
        action="store_true",
        help="Validate all available chunk files in data/ directory",
    )
    parser.add_argument(
        "--master",
        action="store_true",
        help="Validate canonical master JSON (vlearn-app/src/multiphase_course_data.json)",
    )

    args = parser.parse_args()

    print("=" * 70)
    print("VLearn Essay Questions E2E Test & Validation Engine")
    print("=" * 70)

    master_metadata = load_master_metadata(PROJECT_ROOT)
    exit_code = 0

    if args.chunk:
        chunk_path = Path(args.chunk)
        if not chunk_path.is_absolute():
            chunk_path = PROJECT_ROOT / chunk_path
        print(f"\n[Validating Chunk File: {chunk_path.name}]")
        res = validate_chunk_file(chunk_path, master_metadata)
        print(f"Status: {res['status']}")
        print(f"Items: {res.get('total_items', 0)} / Expected: {res.get('expected_items', 'N/A')}")
        print(f"Valid: {res.get('valid_items', 0)}")
        if res.get("missing_items"):
            print(f"Missing ({len(res['missing_items'])}): {res['missing_items'][:10]}")
            exit_code = 1
        if res.get("item_errors"):
            print(f"Errors in {len(res['item_errors'])} items:")
            for qid, errs in list(res["item_errors"].items())[:15]:
                print(f"  [{qid}]: {'; '.join(errs)}")
            exit_code = 1
        if exit_code == 0:
            print(">>> 100% CLEAN & COMPLIANT! <<<")

    elif args.all_chunks:
        data_dir = PROJECT_ROOT / "data"
        chunk_files = sorted(data_dir.glob("essay_block_*.json")) if data_dir.exists() else []
        if not chunk_files:
            print("\nNo chunk files found in data/ directory.")
            return 0
        print(f"\nFound {len(chunk_files)} chunk files in {data_dir}:")
        for cf in chunk_files:
            res = validate_chunk_file(cf, master_metadata)
            status_symbol = "✓" if res["status"] == "CLEAN" else "✗"
            print(f" {status_symbol} {cf.name:25s} [{res['status']}] {res.get('valid_items', 0)}/{res.get('expected_items', '?')} items valid")
            if res["status"] != "CLEAN":
                exit_code = 1

    elif args.master or (not args.chunk and not args.all_chunks):
        print("\n[Validating Master JSON: vlearn-app/src/multiphase_course_data.json]")
        print(f"Indexed metadata for {len(master_metadata)} essay questions across 15 Days.")
        print("Run `pytest tests/test_essay_answers.py` for full test execution.")

    return exit_code


if __name__ == "__main__":
    sys.exit(main())
