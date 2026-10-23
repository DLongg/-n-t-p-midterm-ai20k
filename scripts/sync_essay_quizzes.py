#!/usr/bin/env python3
"""
scripts/sync_essay_quizzes.py

Centralized Cross-File Synchronization Engine for 775 VLearn Essay Questions.
Synchronizes complete model answers, tailored rubrics, explanations, and reference sources
from `vlearn-app/src/multiphase_course_data.json` into:

1. Bank TXT files:
   - `vlearn-app/public/quiz_txt/*_NGAN_HANG_TOAN_DIEN.txt` (15 files)
   - `vlearn-app/public/quiz_txt/DAY*_NGAN_HANG_*_CAU_HOI_TOAN_DIEN.txt` (15 alias files)
   - `downloads/Khoa_4_Phase_1/*/DAY*_NGAN_HANG_*_CAU_HOI_TOAN_DIEN.txt` (15 download mirrors)

2. Master Excel workbooks:
   - `VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx`
   - `VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES_VERIFIED.xlsx`
   - `vlearn-app/public/VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx`
   (Updates Col 9 [Rubric] and Col 10 [Model Answer & References] across 16 bank sheets)

3. Partitioned Excel workbooks:
   - `TONG_HOP_QUIZ_15_NGAY_EXCEL/02_FULL_NGAN_HANG_1679_CAU_CHI_TIET_CHIA_20_CAU/` (84 workbooks)
   (Updates Col 11 [Answer explanation] for all 775 Open-Ended rows)

4. ZIP Archives:
   - `TONG_HOP_QUIZ_15_NGAY_EXCEL.zip`
   - `vlearn-app/public/TONG_HOP_QUIZ_15_NGAY_EXCEL.zip` (100% SHA256 parity, 146 entries)
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def clean_excel_val(val: Any) -> Any:
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub("", val)
    return val


PROJECT_ROOT = Path(__file__).resolve().parent.parent

CANONICAL_MASTER_JSON = PROJECT_ROOT / "vlearn-app" / "src" / "multiphase_course_data.json"

MASTER_XLSX_1 = PROJECT_ROOT / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx"
MASTER_XLSX_2 = PROJECT_ROOT / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES_VERIFIED.xlsx"
MASTER_XLSX_3 = PROJECT_ROOT / "vlearn-app" / "public" / "VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx"

QUIZ_TXT_DIR = PROJECT_ROOT / "vlearn-app" / "public" / "quiz_txt"
DOWNLOADS_K4P1_DIR = PROJECT_ROOT / "downloads" / "Khoa_4_Phase_1"

QUIZIZZ_02_DIR = PROJECT_ROOT / "TONG_HOP_QUIZ_15_NGAY_EXCEL" / "02_FULL_NGAN_HANG_1679_CAU_CHI_TIET_CHIA_20_CAU"

ZIP_MAIN = PROJECT_ROOT / "TONG_HOP_QUIZ_15_NGAY_EXCEL.zip"
ZIP_PUBLIC = PROJECT_ROOT / "vlearn-app" / "public" / "TONG_HOP_QUIZ_15_NGAY_EXCEL.zip"


def load_essay_map() -> Dict[str, Dict[str, Any]]:
    """Loads all 775 essay questions from multiphase_course_data.json."""
    with open(CANONICAL_MASTER_JSON, "r", encoding="utf-8") as f:
        master_data = json.load(f)

    phase1_days = master_data.get("phases", [])[0].get("days", [])[:15]
    essay_map: Dict[str, Dict[str, Any]] = {}
    for day in phase1_days:
        for q in day.get("essay_questions", []):
            qid = q["question_id"]
            essay_map[qid] = {
                "question_id": qid,
                "points": q.get("points"),
                "sample_answer": q.get("sample_answer", ""),
                "rubric": q.get("rubric", ""),
                "explanation": q.get("explanation", ""),
                "reference_source": q.get("reference_source", ""),
                "hints": q.get("hints", []),
            }
    return essay_map


def sync_bank_txt_files(essay_map: Dict[str, Dict[str, Any]]) -> None:
    """Updates 15 bank TXT files and mirrors to aliases and download folders."""
    print("--- [Step 1] Synchronizing Bank TXT Files ---")
    bank_files = sorted(QUIZ_TXT_DIR.glob("*_NGAN_HANG_TOAN_DIEN.txt"))
    alias_files = sorted(QUIZ_TXT_DIR.glob("DAY*_NGAN_HANG_*_CAU_HOI_TOAN_DIEN.txt"))

    assert len(bank_files) == 15, f"Expected 15 bank TXT files, found {len(bank_files)}"
    assert len(alias_files) == 15, f"Expected 15 alias TXT files, found {len(alias_files)}"

    def indent_text(t: str) -> str:
        return "\n".join("  " + line if line.strip() else "" for line in t.splitlines())

    total_txt_updated = 0
    for bf in bank_files:
        content = bf.read_text(encoding="utf-8")
        original_content = content
        file_updated = 0

        # Match each essay question in this bank file
        for qid, qdata in essay_map.items():
            if qid not in content:
                continue

            rubric = qdata["rubric"].strip()
            sample_ans = qdata["sample_answer"].strip()
            exp = qdata["explanation"].strip()
            ref = qdata["reference_source"].strip()

            new_solution_text = (
                f"  CÂU TRẢ LỜI MẪU CHI TIẾT (MODEL ANSWER):\n"
                f"{indent_text(sample_ans)}\n\n"
                f"  PHÂN TÍCH THEO NGUYÊN LÝ THIẾT KẾ:\n"
                f"{indent_text(exp)}\n\n"
                f"  NGUỒN THAM KHẢO & TÀI LIỆU KỸ THUẬT:\n"
                f"  {ref}"
            )

            pat = re.compile(
                rf"(Mã:\s*{re.escape(qid)}[\s\S]*?ĐÁP ÁN ĐÚNG & HƯỚNG DẪN:\s*\n)([\s\S]*?)(\n\s*PHÂN TÍCH LỜI GIẢI / NGUYÊN LÝ THIẾT KẾ TỪ VLEARN:\s*\n)([\s\S]*?)(\n\s*GỢI Ý CỐT LÕI TỪ VLEARN:)",
                re.IGNORECASE,
            )
            m = pat.search(content)
            if m:
                # Replace group 2 with new rubric, group 4 with new solution text
                content = (
                    content[:m.start(2)]
                    + rubric
                    + content[m.end(2):m.start(4)]
                    + new_solution_text
                    + content[m.end(4):]
                )
                file_updated += 1

        if content != original_content:
            bf.write_text(content, encoding="utf-8")
            print(f"Updated {bf.name}: {file_updated} essay questions.")
            total_txt_updated += file_updated

        # Mirror to matching alias in vlearn-app/public/quiz_txt/
        num = bf.name.split("_")[0]
        day_tag = f"DAY{num}_"
        day_tag_alt = f"DAY{int(num)}_"
        matching_aliases = [a for a in alias_files if day_tag in a.name or day_tag_alt in a.name]
        for a in matching_aliases:
            shutil.copy2(bf, a)

        # Mirror to downloads/Khoa_4_Phase_1/*/DAY*...
        dl_matches = list(DOWNLOADS_K4P1_DIR.glob(f"*/DAY{num}_*.txt")) + list(
            DOWNLOADS_K4P1_DIR.glob(f"*/DAY{int(num)}_*.txt")
        )
        for dl_target in dl_matches:
            shutil.copy2(bf, dl_target)

    print(f"Bank TXT synchronization complete: {total_txt_updated} question updates applied and mirrored.")


def sync_master_excel(essay_map: Dict[str, Dict[str, Any]]) -> None:
    """Updates VLEARN_K4P1_TOAN_BO_QUIZ_VA_SLIDES.xlsx across all 16 bank sheets and copies to mirrors."""
    print("--- [Step 2] Synchronizing Master Excel Workbooks ---")
    wb = openpyxl.load_workbook(MASTER_XLSX_1)
    kho_sheets = [s for s in wb.sheetnames if "_Kho_" in s or s == "Day_01_Ngan_Hang_150_Cau"]
    assert len(kho_sheets) == 16, f"Expected 16 bank sheets, found {len(kho_sheets)}"

    total_cells_updated = 0
    for sname in kho_sheets:
        ws = wb[sname]
        sheet_updated = 0
        for r in range(4, ws.max_row + 1):
            qid_val = ws.cell(r, 2).value
            if not qid_val:
                continue
            qid = str(qid_val).strip()
            if qid in essay_map:
                qdata = essay_map[qid]
                # Col 9: Rubric
                ws.cell(r, 9).value = clean_excel_val(qdata["rubric"])
                # Col 10: Model answer, explanation, and reference
                ws.cell(r, 10).value = clean_excel_val(
                    f"CÂU TRẢ LỜI MẪU CHI TIẾT (MODEL ANSWER):\n"
                    f"{qdata['sample_answer']}\n\n"
                    f"PHÂN TÍCH THEO NGUYÊN LÝ THIẾT KẾ:\n"
                    f"{qdata['explanation']}\n\n"
                    f"NGUỒN THAM KHẢO & TÀI LIỆU KỸ THUẬT:\n"
                    f"{qdata['reference_source']}"
                )
                sheet_updated += 1
        print(f"Sheet {sname}: updated {sheet_updated} essay rows.")
        total_cells_updated += sheet_updated

    wb.save(MASTER_XLSX_1)
    wb.close()

    # Mirror to VERIFIED and public
    shutil.copy2(MASTER_XLSX_1, MASTER_XLSX_2)
    shutil.copy2(MASTER_XLSX_1, MASTER_XLSX_3)
    print(f"Master Excel updated ({total_cells_updated} rows across 16 sheets) and mirrored to all 3 paths.")


def sync_partitioned_workbooks(essay_map: Dict[str, Dict[str, Any]]) -> None:
    """Updates Col 11 for all 775 Open-Ended rows in 84 partitioned workbooks."""
    print("--- [Step 3] Synchronizing Partitioned Excel Workbooks (84 files) ---")
    wb_master = openpyxl.load_workbook(MASTER_XLSX_1, data_only=True)
    kho_sheets = [s for s in wb_master.sheetnames if "_Kho_" in s]
    p2_folders = sorted([d for d in QUIZIZZ_02_DIR.iterdir() if d.is_dir()])

    assert len(kho_sheets) == 15
    assert len(p2_folders) == 15

    total_open_ended_updated = 0

    for sname, folder in zip(kho_sheets, p2_folders):
        ws_kho = wb_master[sname]
        kho_qids = [ws_kho.cell(r, 2).value for r in range(4, ws_kho.max_row + 1)]

        current_idx = 0
        folder_updated = 0
        for pf in sorted(folder.glob("*.xlsx")):
            wb_p = openpyxl.load_workbook(pf)
            ws_p = wb_p["Create a Quiz"]
            data_rows = ws_p.max_row - 2

            for r_p in range(1, data_rows + 1):
                qid = str(kho_qids[current_idx]).strip() if current_idx < len(kho_qids) else None
                row_idx = 2 + r_p

                qtype = ws_p.cell(row_idx, 2).value
                if qtype == "Open-Ended" and qid and qid in essay_map:
                    qdata = essay_map[qid]
                    ws_p.cell(row_idx, 11).value = clean_excel_val(
                        f"{qdata['rubric']}\n\n"
                        f"CÂU TRẢ LỜI MẪU CHI TIẾT:\n"
                        f"{qdata['sample_answer']}\n\n"
                        f"PHÂN TÍCH THEO NGUYÊN LÝ THIẾT KẾ:\n"
                        f"{qdata['explanation']}\n\n"
                        f"NGUỒN THAM KHẢO & TÀI LIỆU KỸ THUẬT:\n"
                        f"{qdata['reference_source']}"
                    )
                    folder_updated += 1

                current_idx += 1

            wb_p.save(pf)
            wb_p.close()

        print(f"{folder.name}: updated {folder_updated} Open-Ended essay questions.")
        total_open_ended_updated += folder_updated

    wb_master.close()
    print(f"Partitioned workbooks synchronized: {total_open_ended_updated}/775 Open-Ended rows updated.")
    assert total_open_ended_updated == 775, f"Expected 775 updates, got {total_open_ended_updated}"


def repack_and_mirror_zip() -> None:
    """Repacks TONG_HOP_QUIZ_15_NGAY_EXCEL.zip and mirrors it to vlearn-app/public/."""
    print("--- [Step 4] Repacking and Mirroring ZIP Archives ---")
    source_dir = PROJECT_ROOT / "TONG_HOP_QUIZ_15_NGAY_EXCEL"

    if ZIP_MAIN.exists():
        ZIP_MAIN.unlink()

    with zipfile.ZipFile(ZIP_MAIN, "w", zipfile.ZIP_DEFLATED) as z:
        for file_path in sorted(source_dir.rglob("*")):
            if file_path.is_file():
                rel_path = file_path.relative_to(source_dir).as_posix()
                z.write(file_path, arcname=rel_path)

    shutil.copy2(ZIP_MAIN, ZIP_PUBLIC)

    h1 = hashlib.sha256(ZIP_MAIN.read_bytes()).hexdigest()
    h2 = hashlib.sha256(ZIP_PUBLIC.read_bytes()).hexdigest()
    assert h1 == h2, f"ZIP hash mismatch: {h1} != {h2}"

    with zipfile.ZipFile(ZIP_MAIN, "r") as z:
        entries = [name for name in z.namelist() if not name.endswith("/")]
    assert len(entries) == 146, f"Expected 146 entries in ZIP, found {len(entries)}"

    print(f"ZIP archives synchronized! SHA256: {h1[:16]}... (146 files verified)")


def run_essay_sync() -> None:
    """Executes the complete cross-file synchronization pipeline for essay questions."""
    print("================================================================")
    print("Starting VLearn Essay Questions Cross-File Synchronization")
    print("================================================================")

    essay_map = load_essay_map()
    print(f"Loaded {len(essay_map)} essay questions from master JSON.")
    assert len(essay_map) == 775

    sync_bank_txt_files(essay_map)
    sync_master_excel(essay_map)
    sync_partitioned_workbooks(essay_map)
    repack_and_mirror_zip()

    print("================================================================")
    print("Essay Synchronization Complete: All 4 Targets Fully Synchronized!")
    print("================================================================")


if __name__ == "__main__":
    run_essay_sync()
